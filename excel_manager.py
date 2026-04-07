import pandas as pd
import openpyxl
import os
from telegram import Update
from telegram.ext import ContextTypes

async def convert_plan_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    message = update.message
    
    if not message.document:
        await message.reply_text("⚠️ الرجاء إرسال ملف Excel صحيح.")
        return
        
    try:
        status_msg = await message.reply_text("⏳ جاري تحليل الخطة ونقل البيانات للقالب الجديد... ثواني بس.")
        
        # 1. تحميل الملف القديم من التيليجرام
        file_info = await message.document.get_file()
        old_file_path = f"temp_old_{message.from_user.id}.xlsx"
        await file_info.download_to_drive(custom_path=old_file_path)
            
        # 2. قراءة البيانات
        df_old = pd.read_excel(old_file_path)
        
        # 3. فتح قالب الجودة المعتمد 
        template_path = os.path.join(os.path.dirname(__file__), 'data', 'Curriculum_Plan_v3.xlsx')
        wb_new = openpyxl.load_workbook(template_path)
        ws_new = wb_new.active 
        
        # 🌟 السحر الجديد: دالة الكتابة الآمنة لتخطي مشكلة الخلايا المدمجة 🌟
        def write_safe(ws, r, c, val):
            cell = ws.cell(row=r, column=c)
            if type(cell).__name__ == 'MergedCell':
                # إذا الخلية مدمجة، نبحث عن الخلية الرئيسية في مجموعة الدمج ونكتب فيها
                for merged_range in ws.merged_cells.ranges:
                    if cell.coordinate in merged_range:
                        ws.cell(row=merged_range.min_row, column=merged_range.min_col).value = val
                        return
            else:
                cell.value = val

        # 4. عملية النقل
        # ⚠️ تنبيه: تأكد أن 5 هو رقم أول صف فارغ للبيانات في قالبك. إذا العناوين ماخذة مساحة أكبر، خله 6 أو 7
        start_row = 5 
        
        for index, row in df_old.iterrows():
            current_row = start_row + index
            
            # استخدام الدالة الآمنة بدلاً من الكتابة المباشرة
            write_safe(ws_new, current_row, 1, row.get("رقم الوحدة", ""))
            write_safe(ws_new, current_row, 2, row.get("اسم الوحدة", ""))
            write_safe(ws_new, current_row, 3, row.get("الساعات التدريبية", ""))
            write_safe(ws_new, current_row, 4, row.get("الأهداف التفصيلية", ""))
            write_safe(ws_new, current_row, 5, row.get("موضوعات التدريب", ""))

        # 5. حفظ وإرسال الملف المحدث
        output_filename = f"الخطة_المحدثة_اصدار_3.xlsx"
        wb_new.save(output_filename)
        
        with open(output_filename, 'rb') as doc:
            await message.reply_document(document=doc, caption="🎉 تم الانتهاء بنجاح!\nالخطة المحدثة جاهزة ومطابقة لقالب وكالة الجودة 100%.")
            
        # تنظيف السيرفر من الملفات المؤقتة
        os.remove(old_file_path)
        os.remove(output_filename)
        await status_msg.delete() # مسح رسالة "جاري التحليل" للترتيب
        
        # إنهاء حالة البوت
        context.user_data['waiting_for_plan'] = False
        
    except Exception as e:
        await message.reply_text(f"⚠️ حدث خطأ: {str(e)}\nتأكد من مطابقة أسماء الأعمدة في ملفك القديم.")
        context.user_data['waiting_for_plan'] = False
